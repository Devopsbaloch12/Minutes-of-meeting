"""Minutes as PDF and Word documents.

Both formats are built from the sanitised HTML the meeting page shows, through
one small parser, so a download always matches what is on screen. Only what
model-written Markdown produces is handled: headings, paragraphs, bold and
italic, nested lists, tables, rules and code.
"""
import io
import os
import re
import unicodedata
from dataclasses import dataclass, field
from html.parser import HTMLParser

from app.utils.text import render_markdown

_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}
_WHITESPACE = re.compile(r"\s+")


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str  # heading | paragraph | item | table | rule | pre
    runs: list[Run] = field(default_factory=list)
    level: int = 0  # heading level, or list depth starting at 1
    marker: str = ""  # "" for a bullet, "3." for a numbered item
    rows: list[list[str]] = field(default_factory=list)


class _Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.current: Block | None = None
        self.lists: list[list] = []  # [tag, next number]
        self.bold = self.italic = self.code = 0
        self.pre = False
        self.rows: list[list[str]] | None = None
        self.cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag in _HEADINGS:
            self._open("heading", level=_HEADINGS[tag])
        elif tag == "p" and self.cell is None:
            if not (self.current and self.current.kind == "item"):
                self._open("paragraph")
        elif tag in ("ul", "ol"):
            self._flush()  # a nested list ends the text of the item that holds it
            start = dict(attrs).get("start") or "1"
            self.lists.append([tag, int(start) if start.isdigit() else 1])
        elif tag == "li":
            depth = max(1, len(self.lists))
            marker = ""
            if self.lists and self.lists[-1][0] == "ol":
                marker = f"{self.lists[-1][1]}."
                self.lists[-1][1] += 1
            self._open("item", level=depth, marker=marker)
        elif tag in ("strong", "b"):
            self.bold += 1
        elif tag in ("em", "i"):
            self.italic += 1
        elif tag == "code" and not self.pre:
            self.code += 1
        elif tag == "br":
            self._text("\n")
        elif tag == "table":
            self._flush()
            self.rows = []
        elif tag == "tr" and self.rows is not None:
            self.rows.append([])
        elif tag in ("th", "td") and self.rows is not None:
            self.cell = []
        elif tag == "pre":
            self._open("pre")
            self.pre = True
        elif tag == "hr":
            self._flush()
            self.blocks.append(Block("rule"))

    def handle_endtag(self, tag):
        if tag in _HEADINGS or tag in ("p", "li", "pre"):
            if tag == "pre":
                self.pre = False
            if not (tag == "p" and self.current and self.current.kind == "item"):
                self._flush()
        elif tag in ("ul", "ol") and self.lists:
            self._flush()
            self.lists.pop()
        elif tag in ("strong", "b"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("em", "i"):
            self.italic = max(0, self.italic - 1)
        elif tag == "code" and not self.pre:
            self.code = max(0, self.code - 1)
        elif tag in ("th", "td") and self.cell is not None and self.rows:
            self.rows[-1].append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "table" and self.rows is not None:
            rows = [row for row in self.rows if row]
            if rows:
                width = max(len(row) for row in rows)
                # Model tables are not always rectangular; both writers need them to be.
                self.blocks.append(Block("table", rows=[row + [""] * (width - len(row)) for row in rows]))
            self.rows = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)
        elif self.pre:
            self._text(data)
        elif data.strip() or (self.current and self.current.runs):
            # Newlines here are HTML source formatting; <br> is the real line break.
            # Edge spaces stay: they separate words from a following <strong>.
            self._text(_WHITESPACE.sub(" ", data))

    def _open(self, kind, level=0, marker=""):
        self._flush()
        self.current = Block(kind, level=level, marker=marker)

    def _text(self, text):
        if self.current is None:
            self.current = Block("paragraph")
        self.current.runs.append(Run(text, self.bold > 0, self.italic > 0, self.code > 0))

    def _flush(self):
        block, self.current = self.current, None
        if block is None:
            return
        while block.runs and not block.runs[-1].text.strip():
            block.runs.pop()
        if block.runs:
            block.runs[0].text = block.runs[0].text.lstrip()
            self.blocks.append(block)


def parse_minutes(markdown_text: str) -> list[Block]:
    parser = _Parser()
    parser.feed(render_markdown(markdown_text))
    parser.close()
    parser._flush()
    return parser.blocks


# ── Word ──────────────────────────────────────────────────────────────────


def minutes_to_docx(markdown_text: str, title: str, details: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.add_heading(title, level=0)
    if details:
        run = document.add_paragraph().add_run(details)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    blocks = parse_minutes(markdown_text)
    top = min((b.level for b in blocks if b.kind == "heading"), default=1)
    for block in blocks:
        if block.kind == "heading":
            document.add_heading(_plain(block), level=min(4, block.level - top + 1))
        elif block.kind == "paragraph":
            _add_runs(document.add_paragraph(), block.runs)
        elif block.kind == "item":
            if block.marker:
                # Word's "List Number" style shares one counter across the whole document,
                # so Decisions would continue from the Agenda; write the number instead.
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(18 * block.level)
                paragraph.paragraph_format.first_line_indent = Pt(-14)
                _add_runs(paragraph, [Run(f"{block.marker} ")] + block.runs)
            else:
                style = "List Bullet" if block.level == 1 else f"List Bullet {min(block.level, 3)}"
                _add_runs(document.add_paragraph(style=style), block.runs)
        elif block.kind == "table":
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(block.rows):
                for c, text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    cell.paragraphs[0].add_run(text).bold = r == 0
            document.add_paragraph()
        elif block.kind == "pre":
            paragraph = document.add_paragraph()
            _add_runs(paragraph, [Run(_plain(block), code=True)])
        elif block.kind == "rule":
            document.add_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_runs(paragraph, runs: list[Run]) -> None:
    for item in runs:
        for index, line in enumerate(item.text.split("\n")):
            if index:
                paragraph.add_run().add_break()
            if not line:
                continue
            run = paragraph.add_run(line)
            run.bold = item.bold or None
            run.italic = item.italic or None
            if item.code:
                run.font.name = "Consolas"


# ── PDF ───────────────────────────────────────────────────────────────────

# Core PDF fonts only cover Latin-1, and model output is full of €, ≈, curly
# quotes and non-breaking hyphens. The first family found on the machine is
# used; the Docker image ships DejaVu (fonts-dejavu-core) for exactly this.
_WINDOWS_FONTS = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
_FONT_FAMILIES = (
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"),
    tuple(os.path.join(_WINDOWS_FONTS, name) for name in ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
    tuple(f"/System/Library/Fonts/Supplemental/{name}.ttf"
          for name in ("Arial", "Arial Bold", "Arial Italic", "Arial Bold Italic")),
)
# Typography models emit that common fonts, Arial included, have no glyph for:
# gpt-oss writes "user‑friendly" with U+2011, which printed as "userfriendly".
_UNPRINTABLE = str.maketrans({"‑": "-", " ": " ", " ": " "})
_LATIN1_SUBSTITUTES = str.maketrans({
    "‐": "-", "–": "-", "—": "-", "‘": "'", "’": "'", "“": '"',
    "”": '"', "•": "-", "…": "...", "€": "EUR", "≈": "~", "≤": "<=",
    "≥": ">=", "→": "->",
})


def minutes_to_pdf(markdown_text: str, title: str, details: str = "") -> bytes:
    from fpdf import FPDF, FontFace

    class Document(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font(family, "", 8)
            self.set_text_color(150)
            self.cell(0, 5, clean(f"{title} - page {self.page_no()}"), align="C")

    pdf = Document(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    family = _register_font(pdf)

    def clean(text: str) -> str:
        text = text.translate(_UNPRINTABLE)
        if family != "helvetica":
            return text
        text = unicodedata.normalize("NFKC", text).translate(_LATIN1_SUBSTITUTES)
        return text.encode("latin-1", "replace").decode("latin-1")

    pdf.add_page()
    pdf.set_font(family, "B", 18)
    pdf.multi_cell(0, 9, clean(title), new_x="LMARGIN", new_y="NEXT")
    if details:
        pdf.set_font(family, "", 9)
        pdf.set_text_color(107, 114, 128)
        pdf.multi_cell(0, 5, clean(details), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0)
    pdf.ln(3)

    body = 10
    blocks = parse_minutes(markdown_text)
    top = min((b.level for b in blocks if b.kind == "heading"), default=1)
    for block in blocks:
        if block.kind == "heading":
            size = {1: 14, 2: 12}.get(block.level - top + 1, 11)
            pdf.ln(3)
            pdf.set_font(family, "B", size)
            pdf.multi_cell(0, size * 0.55, clean(_plain(block)), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif block.kind in ("paragraph", "pre"):
            _write_runs(pdf, family, body, block.runs, clean)
            pdf.ln(2)
        elif block.kind == "item":
            indent = 5 * block.level
            marker = block.marker or ("•" if family != "helvetica" else "-")
            pdf.set_x(pdf.l_margin + indent - 4)
            pdf.set_font(family, "", body)
            pdf.cell(4 if not block.marker else 6, body * 0.5, clean(marker))
            left = pdf.l_margin
            pdf.set_left_margin(left + indent + (2 if block.marker else 0))
            _write_runs(pdf, family, body, block.runs, clean)
            pdf.set_left_margin(left)
            pdf.ln(0.5)
        elif block.kind == "table":
            pdf.set_font(family, "", body - 1)
            # Equal columns squeeze the Action text into a sliver beside a short Owner.
            widths = [min(60, max(len(row[c]) for row in block.rows)) + 6 for c in range(len(block.rows[0]))]
            with pdf.table(
                col_widths=widths, text_align="LEFT", line_height=(body - 1) * 0.5, padding=1.5,
                headings_style=FontFace(emphasis="BOLD", fill_color=(243, 244, 246)),
            ) as table:
                for row in block.rows:
                    cells = table.row()
                    for text in row:
                        cells.cell(clean(text))
            pdf.ln(3)
        elif block.kind == "rule":
            pdf.ln(2)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(3)

    return bytes(pdf.output())


def _register_font(pdf) -> str:
    for regular, bold, italic, bold_italic in _FONT_FAMILIES:
        if not os.path.exists(regular):
            continue
        bold = bold if os.path.exists(bold) else regular
        # fonts-dejavu-core has no oblique faces; upright beats failing the export.
        italic = italic if os.path.exists(italic) else regular
        bold_italic = bold_italic if os.path.exists(bold_italic) else bold
        for style, path in (("", regular), ("B", bold), ("I", italic), ("BI", bold_italic)):
            pdf.add_font("body", style, path)
        return "body"
    return "helvetica"


def _write_runs(pdf, family: str, size: float, runs: list[Run], clean) -> None:
    """write() wraps at the left margin, so callers indent by moving it."""
    for run in runs:
        style = ("B" if run.bold else "") + ("I" if run.italic else "")
        pdf.set_font(family, style, size - 1 if run.code else size)
        pdf.write(size * 0.5, clean(run.text))
    pdf.ln(size * 0.5)


def _plain(block: Block) -> str:
    return "".join(run.text for run in block.runs).strip()
