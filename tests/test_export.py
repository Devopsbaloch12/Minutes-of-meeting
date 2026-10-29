import io

from docx import Document

from app.utils import export

MINUTES = """## Summary

Revenue rose **10%** and the team agreed to *ship* in May.

## Agenda

1. Budget
2. Launch date

## Action Items

| # | Action | Owner | Due date |
| --- | --- | --- | --- |
| 1 | Send the deck | Ana | Friday |
| 2 | Book the venue | Unassigned |

## Next Steps

- Review on Monday
    - Ana presents
"""


def test_minutes_are_parsed_into_document_blocks():
    blocks = export.parse_minutes(MINUTES)
    kinds = [(b.kind, b.level, b.marker) for b in blocks]

    assert kinds == [
        ("heading", 2, ""), ("paragraph", 0, ""),
        ("heading", 2, ""), ("item", 1, "1."), ("item", 1, "2."),
        ("heading", 2, ""), ("table", 0, ""),
        ("heading", 2, ""), ("item", 1, ""), ("item", 2, ""),
    ]
    summary = blocks[1].runs
    assert "".join(r.text for r in summary) == "Revenue rose 10% and the team agreed to ship in May."
    assert [r.text for r in summary if r.bold] == ["10%"] and [r.text for r in summary if r.italic] == ["ship"]


def test_ragged_model_tables_are_padded_to_a_rectangle():
    table = next(b for b in export.parse_minutes(MINUTES) if b.kind == "table")
    assert table.rows[0] == ["#", "Action", "Owner", "Due date"]
    assert table.rows[2] == ["2", "Book the venue", "Unassigned", ""]


def test_word_export_carries_headings_tables_and_numbers():
    document = Document(io.BytesIO(export.minutes_to_docx(MINUTES, "Weekly sync", "27 Oct 2026")))
    text = [p.text for p in document.paragraphs]

    assert text[0] == "Weekly sync" and "27 Oct 2026" in text
    assert "Summary" in text and "Action Items" in text
    assert "1. Budget" in text and "2. Launch date" in text
    cells = [[c.text for c in row.cells] for row in document.tables[0].rows]
    assert cells[1] == ["1", "Send the deck", "Ana", "Friday"]


def test_pdf_export_is_a_pdf_document():
    pdf = export.minutes_to_pdf(MINUTES, "Weekly sync", "27 Oct 2026")
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF")


def test_pdf_export_survives_without_a_unicode_font(monkeypatch):
    """Core PDF fonts are Latin-1 only; model output is not."""
    monkeypatch.setattr(export, "_FONT_FAMILIES", ())
    text = "## Summary\n\nCost ≤ €12.50 – Laura’s “user‑friendly” remote → 中文"
    assert export.minutes_to_pdf(text, "Kick‑off – €").startswith(b"%PDF")
